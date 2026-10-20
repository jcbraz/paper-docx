"""Compare must distinguish serialization differences from document edits."""

import datetime as dt
import zipfile
from pathlib import Path

import pytest
from lxml import etree

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.package import compare

FROZEN = dt.datetime(2026, 7, 10, 12, tzinfo=dt.timezone.utc)
MC_IGNORABLE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"


@pytest.mark.parametrize("revised_text", ["Alpha content", "Alpha revised content"])
def it_compares_reordered_story_attributes_and_ignorable_prefixes(
    tmp_path: Path, revised_text: str
):
    paths: list[Path] = []
    for name, text in (("original", "Alpha content"), ("revised", revised_text)):
        document = docx.Document()
        paragraph = document.add_paragraph(text)
        spacing = OxmlElement("w:spacing")
        attributes = [("w:before", "20"), ("w:after", "120")]
        if name == "revised":
            attributes.reverse()
        for key, value in attributes:
            spacing.set(qn(key), value)
        paragraph._p.get_or_add_pPr().append(spacing)
        # Exercise the whole-story contract as well as paragraph preflight.
        footer = document.sections[0].footer.paragraphs[0]
        footer.text = "Footer"
        footer._p.get_or_add_pPr().append(etree.fromstring(etree.tostring(spacing)))
        path = tmp_path / f"{name}.docx"
        document.save(str(path))
        paths.append(path)
    _reverse_ignorable_prefixes(paths[1])
    before = [path.read_bytes() for path in paths]

    result = compare(*paths, author="Reviewer", date=FROZEN)
    output = tmp_path / "redline.docx"
    result.document.save(str(output))
    with zipfile.ZipFile(paths[0]) as source, zipfile.ZipFile(output) as redline:
        assert set(redline.namelist()) == set(source.namelist())
        for part in source.namelist():
            if part != "word/document.xml":
                assert redline.read(part) == source.read(part), part
    accepted = docx.Document(str(output))
    accepted.revisions.accept_all()
    assert [p.text for p in accepted.paragraphs] == [revised_text]
    rejected = docx.Document(str(output))
    rejected.revisions.reject_all()
    assert [p.text for p in rejected.paragraphs] == ["Alpha content"]
    assert [path.read_bytes() for path in paths] == before


def it_merges_fragmented_runs_with_equivalent_property_attributes(tmp_path: Path):
    paths: list[Path] = []
    for name, fragments in (("original", ["Alpha content"]), ("revised", ["Alpha ", "content"])):
        document = docx.Document()
        paragraph = document.add_paragraph()
        for index, text in enumerate(fragments):
            run = paragraph.add_run(text)
            fonts = OxmlElement("w:rFonts")
            attributes = [("w:ascii", "Arial"), ("w:hAnsi", "Arial")]
            if index:
                attributes.reverse()
            for key, value in attributes:
                fonts.set(qn(key), value)
            run._r.get_or_add_rPr().append(fonts)
        path = tmp_path / f"{name}.docx"
        document.save(str(path))
        paths.append(path)
    result = compare(*paths, author="Reviewer", date=FROZEN)
    output = tmp_path / "redline.docx"
    result.document.save(str(output))
    reopened = docx.Document(str(output))
    assert len(reopened.revisions) == 0
    assert reopened.paragraphs[0].text == "Alpha content"


def _reverse_ignorable_prefixes(path: Path):
    with zipfile.ZipFile(path) as source:
        parts = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w") as target:
        for info, data in parts:
            if info.filename.endswith(".xml"):
                root = etree.fromstring(data)
                for element in root.iter():
                    value = element.get(MC_IGNORABLE)
                    if value:
                        element.set(MC_IGNORABLE, " ".join(reversed(value.split())))
                data = etree.tostring(root)
            target.writestr(info, data)
